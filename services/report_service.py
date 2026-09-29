"""
services/report_service.py
Generación de reportes institucionales en PDF (ReportLab) y Excel (openpyxl).
"""

from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from database.connection import get_connection

class ReportService:

    @staticmethod
    def _obtener_datos_prestamos(fecha_desde=None, fecha_hasta=None, estado=None):
        conn = get_connection()
        cur = conn.cursor()

        query = """
            SELECT p.id, p.fecha, p.hora_inicio, p.hora_fin, p.observaciones, p.estado, p.fecha_devolucion,
                   f.nombre as profesor, c.nombre as curso
            FROM prestamos p
            JOIN funcionarios f ON p.funcionario_id = f.id
            LEFT JOIN cursos c ON p.curso_id = c.id
            WHERE 1=1
        """
        params = []

        if fecha_desde:
            query += " AND p.fecha >= ?"
            params.append(fecha_desde)
        if fecha_hasta:
            query += " AND p.fecha <= ?"
            params.append(fecha_hasta)
        if estado and estado != "todos":
            query += " AND p.estado = ?"
            params.append(estado)

        query += " ORDER BY p.fecha DESC, p.hora_inicio DESC;"
        cur.execute(query, params)
        prestamos = cur.fetchall()

        # Adjuntar lista de recursos
        resultado = []
        for p in prestamos:
            cur.execute("""
                SELECT r.nombre, pd.cantidad
                FROM prestamo_detalles pd
                JOIN recursos r ON pd.recurso_id = r.id
                WHERE pd.prestamo_id = ?;
            """, (p["id"],))
            detalles = cur.fetchall()
            recursos_str = ", ".join([f"{d['cantidad']} {d['nombre']}" if d["cantidad"] > 1 else d["nombre"] for d in detalles])

            resultado.append({
                "fecha": p["fecha"],
                "horario": f"{p['hora_inicio']} - {p['hora_fin']}",
                "profesor": p["profesor"],
                "curso": p["curso"] or "--",
                "recursos": recursos_str,
                "observaciones": p["observaciones"] or "",
                "estado": p["estado"].capitalize()
            })

        conn.close()
        return resultado

    @staticmethod
    def exportar_excel(ruta_destino: str, fecha_desde=None, fecha_hasta=None, estado=None) -> bool:
        datos = ReportService._obtener_datos_prestamos(fecha_desde, fecha_hasta, estado)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial Prestamos"

        # Título y membrete
        ws.merge_cells("A1:G1")
        ws["A1"] = "ESCUELA ARAUCANÍA 510 - REGISTRO DE PRÉSTAMOS DE SALA Y RECURSOS"
        ws["A1"].font = Font(size=14, bold=True, color="1E1E24")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} | Sistema AraucaníaStock by KoaLink"
        ws["A2"].font = Font(size=10, italic=True, color="555555")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        # Encabezados
        headers = ["Fecha", "Horario", "Solicitante", "Curso", "Recurso(s)", "Observaciones", "Estado"]
        ws.append([]) # Fila 3 vacía
        ws.append(headers) # Fila 4

        header_fill = PatternFill(start_color="F4D06F", end_color="F4D06F", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="1E1E24")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        for col_num in range(1, 8):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Cargar datos
        for d in datos:
            row = [d["fecha"], d["horario"], d["profesor"], d["curso"], d["recursos"], d["observaciones"], d["estado"]]
            ws.append(row)

        # Ajuste de anchos y bordes
        for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=7):
            for cell in row:
                cell.border = thin_border
                cell.font = Font(name="Segoe UI", size=10)

        anchos = [14, 16, 25, 12, 35, 30, 14]
        for idx, ancho in enumerate(anchos, start=1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws.column_dimensions[col_letter].width = ancho

        wb.save(ruta_destino)
        return True

    @staticmethod
    def exportar_pdf(ruta_destino: str, fecha_desde=None, fecha_hasta=None, estado=None) -> bool:
        datos = ReportService._obtener_datos_prestamos(fecha_desde, fecha_hasta, estado)

        doc = SimpleDocTemplate(
            ruta_destino,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        estilo_titulo = ParagraphStyle(
            name='Titulo',
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#1E1E24"),
            alignment=1
        )
        estilo_sub = ParagraphStyle(
            name='Subtitulo',
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#555555"),
            alignment=1
        )
        estilo_celda = ParagraphStyle(
            name='Celda',
            fontName='Helvetica',
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#1E1E24")
        )

        story = []

        story.append(Paragraph("ESCUELA ARAUCANÍA 510 LABRANZA", estilo_titulo))
        story.append(Paragraph("Informe Oficial de Uso y Préstamos de Sala Multiuso y Recursos", estilo_sub))
        story.append(Paragraph(f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')} — KoaLink Solutions", estilo_sub))
        story.append(Spacer(1, 15))

        # Tabla de Datos
        data_table = [["Fecha", "Horario", "Solicitante", "Curso", "Recurso(s)", "Estado"]]

        for d in datos:
            data_table.append([
                Paragraph(d["fecha"], estilo_celda),
                Paragraph(d["horario"], estilo_celda),
                Paragraph(d["profesor"], estilo_celda),
                Paragraph(d["curso"], estilo_celda),
                Paragraph(d["recursos"], estilo_celda),
                Paragraph(d["estado"], estilo_celda)
            ])

        t = Table(data_table, colWidths=[65, 80, 110, 50, 165, 70])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F4D06F")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor("#1E1E24")),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))

        story.append(t)
        doc.build(story)
        return True